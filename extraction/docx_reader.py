"""
docx_reader.py

Extracts raw text (paragraphs + tables) from a .docx file, in the same
spirit as your existing pdf_reader.py (digital PDFs) and ocr_reader.py
(scanned PDFs).

Design note:
    This module ONLY extracts raw text. It does NOT try to guess which
    field is which (that stays field_extractor.py's job). Keeping this
    separation means field_extractor.py can treat text from a DOCX the
    exact same way it treats text pulled out of a PDF or an OCR pass -
    one extraction brain, three input formats.

Dependency: pip install python-docx --break-system-packages
"""

from docx import Document
from docx.table import Table
from docx.text.paragraph import Paragraph


def read_docx(file_path: str) -> str:
    """
    Extract all readable text from a .docx file: paragraphs AND tables,
    in the order they appear in the document.

    Government certificates are often laid out as tables (label | value),
    e.g. "GSTIN | 27ABCPL1234F1Z5", so table cells are flattened into
    "Label: Value" style lines - this keeps them regex-friendly for
    field_extractor.py, which already expects that kind of pattern from
    scanned certificates.

    Args:
        file_path: path to the .docx file on disk.

    Returns:
        A single string of extracted text, paragraphs and table rows
        separated by newlines - same shape as what pdf_reader.py /
        ocr_reader.py hand back today.

    Raises:
        ValueError: if the file can't be opened as a valid .docx
                    (e.g. it's actually a .doc, or corrupted).
    """
    try:
        doc = Document(file_path)
    except Exception as exc:
        raise ValueError(
            f"Could not read '{file_path}' as a .docx file. "
            f"If this is a legacy .doc file, convert it to .docx first. "
            f"Original error: {exc}"
        ) from exc

    lines = []

    # Walk the document body in order so tables that sit between
    # paragraphs don't get shuffled to the end.
    for block in _iter_block_items(doc):
        if isinstance(block, Paragraph):
            text = block.text.strip()
            if text:
                lines.append(text)
        elif isinstance(block, Table):
            lines.extend(_extract_table_text(block))

    return "\n".join(lines)


def _extract_table_text(table: Table) -> list:
    """
    Flatten a docx table into 'Label: Value' style lines.

    Most certificate tables are 2 columns (label, value) like the GST
    certificate screenshot - so this specifically special-cases 2-column
    rows into "Label: Value". Wider tables fall back to pipe-separated
    cells so nothing is silently dropped.
    """
    rows_text = []
    for row in table.rows:
        cells = [cell.text.strip() for cell in row.cells]
        cells = [c for c in cells if c]  # drop empty cells
        if not cells:
            continue
        if len(cells) == 2:
            rows_text.append(f"{cells[0]}: {cells[1]}")
        else:
            rows_text.append(" | ".join(cells))
    return rows_text


def _iter_block_items(doc: Document):
    """
    Yield paragraphs and tables in the order they actually appear in
    the document body. python-docx doesn't expose this directly - it
    gives you doc.paragraphs and doc.tables as two separate flat lists,
    which loses ordering when a table sits between two paragraphs.
    """
    from docx.oxml.ns import qn

    body = doc.element.body
    for child in body.iterchildren():
        if child.tag == qn("w:p"):
            yield Paragraph(child, doc)
        elif child.tag == qn("w:tbl"):
            yield Table(child, doc)


if __name__ == "__main__":
    # Quick manual test:
    #   python docx_reader.py path/to/certificate.docx
    import sys

    if len(sys.argv) != 2:
        print("Usage: python docx_reader.py <path_to_docx>")
        sys.exit(1)

    extracted_text = read_docx(sys.argv[1])
    print("----- Extracted text -----")
    print(extracted_text)
